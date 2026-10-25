import json
import contextlib
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

import main
from src.config import PROJECT, select_detector, select_llm
from src.data import prepare_dataset, read_json
from src.llm import final_text, generate
from src.products import from_qa_csv
from src.review import choose_with_llm, create_excel, rank_rows, read_review, save_choices, select_training_rows


def candidates(products=5):
    return [{"product_id": f"p{i}", "product_description": f"Product {i}",
             "candidate_id": f"p{i}_{j}", "question": f"Funny sentence {i} {j}."}
            for i in range(products) for j in range(3)]


class WorkflowTests(unittest.TestCase):
    def test_seed_is_optional_and_paths_relocate(self):
        default = select_detector("deberta")
        selected = select_detector("deberta", 44, PROJECT / "models/detector")
        self.assertEqual(default["seed"], 42)
        self.assertIn("qproduct_seed44", selected["checkpoint"])
        self.assertEqual(select_detector("roberta-qonly")["input_mode"], "text")
        self.assertEqual(select_detector("lstm")["seed"], None)
        self.assertTrue(select_detector("cnn")["checkpoint"].endswith("cnn_fasttext_question_only.pt"))

    def test_llm_registry_and_adapter_backend(self):
        self.assertEqual(select_llm("deepseek")["backend"], "transformers")
        self.assertEqual(select_llm("qwen", "ollama")["model"], "qwen2.5:7b")
        with self.assertRaises(ValueError):
            select_llm("qwen", "ollama", adapter="models/llm/adapters/example")

    def test_three_sentence_generation_and_reasoning(self):
        class Model:
            config = {"name": "test", "model": "test-model"}

            def answer(self, prompt, sample=False):
                return "Your coffee works overtime.\nYour mug needs a vacation.\nOffice politics in a cup."

        rows = generate(Model(), [{"product_id": "mug", "product_description": "Coffee mug"}], 3)
        self.assertEqual(len(rows), 3)
        self.assertEqual(len({r["candidate_id"] for r in rows}), 3)
        self.assertEqual(final_text("<think>private reasoning</think>Final sentence."), "Final sentence.")
        self.assertEqual(final_text("reasoning continued from template</think>Final sentence."), "Final sentence.")
        with self.assertRaises(ValueError):
            final_text("<think>unfinished reasoning")

    def test_ranks_are_per_product(self):
        ranked = rank_rows(candidates(2), [.1, .9, .4, .8, .2, .3])
        self.assertEqual([r["rank"] for r in ranked], [1, 2, 3, 1, 2, 3])
        self.assertEqual([ranked[0]["candidate_id"], ranked[3]["candidate_id"]], ["p0_1", "p1_0"])

    def test_ai_one_of_three_preserves_human_selection(self):
        rows = rank_rows(candidates(2), [.3] * 6)
        rows[0]["selected"] = 1
        class Model:
            config = {"name": "reviewer"}

            def answer(self, prompt):
                return json.dumps({"candidate_id": "p1_2", "reason": "Clear comic twist"})

        reviewed = choose_with_llm(rows, Model())
        selected = select_training_rows(reviewed)
        self.assertEqual({r["candidate_id"] for r in selected}, {"p0_0", "p1_2"})
        self.assertEqual(rows[3]["selected"], None)

    def test_selection_requires_exactly_one_per_product(self):
        rows = rank_rows(candidates(1), [.3] * 3)
        with self.assertRaises(ValueError):
            select_training_rows(rows)
        rows[0]["selected"] = rows[1]["selected"] = 1
        with self.assertRaises(ValueError):
            select_training_rows(rows)

    def test_excel_roundtrip_and_original_columns(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory(dir=PROJECT) as tmp:
            source, output = Path(tmp) / "ranked.xlsx", Path(tmp) / "reviewed.xlsx"
            rows = rank_rows(candidates(3), [.3] * 9)
            rows[0]["tagline"] = "=This remains text"
            create_excel(rows, source)
            workbook = load_workbook(source)
            self.assertEqual(workbook.sheetnames, ["Human_review", "All_scores"])
            self.assertEqual(workbook["Human_review"].max_column, 8)
            self.assertEqual(workbook["Human_review"]["D2"].data_type, "s")
            self.assertEqual(workbook["Human_review"]["G2"].value, None)
            workbook.close()
            reviewed = read_review(source)
            for row in reviewed:
                row["selected"] = int(row["rank"] == 1)
            save_choices(source, output, reviewed)
            self.assertEqual(len(select_training_rows(read_review(output))), 3)
            workbook = load_workbook(output)
            self.assertEqual(workbook["All_scores"].max_row, 10)
            workbook.close()
            with self.assertRaises(FileExistsError):
                create_excel(rows, source)

    def test_original_split_ratio_without_product_leakage(self):
        with tempfile.TemporaryDirectory(dir=PROJECT) as tmp:
            dest = Path(tmp) / "lora"
            rows = [r for r in candidates(50) if r["candidate_id"].endswith("_0")]
            self.assertEqual(prepare_dataset(rows, dest), {"train": 35, "validation": 5, "test": 10})
            membership = read_json(dest / "split_manifest.json")["products"]
            self.assertEqual(len(membership), 50)
            self.assertEqual(len(set(membership)), 50)

    def test_qa_products(self):
        with tempfile.TemporaryDirectory(dir=PROJECT) as tmp:
            source = Path(tmp) / "train.csv"
            source.write_text("product_group,product_description,label\np1,Mug,0\np1,Mug,1\np2,Fan,1\n", encoding="utf-8")
            products = from_qa_csv(source)
            self.assertEqual(len(products), 2)
            self.assertNotIn("label", products[0])

    def test_rank_command_does_not_generate_or_train(self):
        with tempfile.TemporaryDirectory(dir=PROJECT) as tmp:
            source, output = Path(tmp) / "sentences.json", Path(tmp) / "ranking.xlsx"
            source.write_text(json.dumps(candidates(3)), encoding="utf-8")
            factory = lambda *args: types.SimpleNamespace(score=lambda rows: [.6] * len(rows))
            with patch.object(main, "preflight", side_effect=lambda config: config), \
                 patch("src.detector.create_detector", factory):
                main.main(["rank", "--input", str(source), "--output", str(output)])
            self.assertTrue(output.is_file())
            self.assertEqual(len(read_review(output)), 9)
            with self.assertRaisesRegex(ValueError, "select exactly one"):
                main.main(["prepare", "--input", str(output), "--output", str(Path(tmp) / "lora")])

    def test_base_adapter_comparison_uses_same_heldout_prompt(self):
        from src.testing import compare
        from openpyxl import load_workbook
        class FakeLLM:
            def __init__(self, *args, **kwargs):
                self.disabled = False
                self.model = self

            @contextlib.contextmanager
            def disable_adapter(self):
                self.disabled = True
                try:
                    yield
                finally:
                    self.disabled = False

            def answer(self, prompt):
                return ("base: " if self.disabled else "tuned: ") + prompt

        with tempfile.TemporaryDirectory(dir=PROJECT) as tmp:
            test = Path(tmp) / "test.jsonl"
            test.write_text(json.dumps({"messages": [{"role": "user", "content": "held-out product"},
                                                     {"role": "assistant", "content": "reference sentence"}]}) + "\n", encoding="utf-8")
            output = Path(tmp) / "test.xlsx"
            with patch("src.llm.LanguageModel", FakeLLM), patch("src.artifacts.check_weights"):
                compare({"model": "test-model"}, Path(tmp) / "adapter", test, output)
            workbook = load_workbook(output)
            row = list(workbook.active.values)[1]
            self.assertEqual(row, ("held-out product", "reference sentence", "base: held-out product", "tuned: held-out product"))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
